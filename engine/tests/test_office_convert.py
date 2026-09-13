from __future__ import annotations

import shutil

import docx
import openpyxl
import pytest
from pypdf import PdfReader

from engine.jobs.contract import JobRequest
from engine.tools.excel_to_pdf import excel_to_pdf
from engine.tools.html_to_pdf import html_to_pdf
from engine.tools.ppt_to_pdf import ppt_to_pdf
from engine.tools.word_to_pdf import word_to_pdf

pytestmark = pytest.mark.skipif(shutil.which("soffice") is None, reason="LibreOffice (soffice) not installed")


def _request(tool: str, **overrides) -> JobRequest:
    base = dict(job_id="job-1", tool=tool, inputs=[], options={}, output_dir="/tmp", overwrite_original=False)
    base.update(overrides)
    return JobRequest(**base)


def test_html_to_pdf_converts_real_content(tmp_path, output_dir):
    html_path = tmp_path / "page.html"
    html_path.write_text("<html><body><h1>Hello</h1><p>World</p></body></html>", encoding="utf-8")

    result = html_to_pdf(_request("html-to-pdf", inputs=[str(html_path)], output_dir=output_dir))

    assert result.status == "success"
    text = PdfReader(result.outputs[0]).pages[0].extract_text()
    assert "Hello" in text
    assert "World" in text
    assert result.warnings


def test_word_to_pdf_rejects_wrong_extension(tmp_path, output_dir):
    bogus = tmp_path / "notes.txt"
    bogus.write_text("hello")

    result = word_to_pdf(_request("word-to-pdf", inputs=[str(bogus)], output_dir=output_dir))

    assert result.status == "error"
    assert result.error.code == "INPUT_TYPE_INVALID"


def test_word_to_pdf_converts_real_content(tmp_path, output_dir):
    doc = docx.Document()
    doc.add_paragraph("A real paragraph of text for conversion testing.")
    docx_path = tmp_path / "doc.docx"
    doc.save(str(docx_path))

    result = word_to_pdf(_request("word-to-pdf", inputs=[str(docx_path)], output_dir=output_dir))

    assert result.status == "success"
    text = PdfReader(result.outputs[0]).pages[0].extract_text()
    assert "A real paragraph" in text


def test_excel_to_pdf_converts_real_content(tmp_path, output_dir):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "HeaderCell"
    xlsx_path = tmp_path / "sheet.xlsx"
    wb.save(str(xlsx_path))

    result = excel_to_pdf(_request("excel-to-pdf", inputs=[str(xlsx_path)], output_dir=output_dir))

    assert result.status == "success"
    text = PdfReader(result.outputs[0]).pages[0].extract_text()
    assert "HeaderCell" in text


def test_ppt_to_pdf_rejects_missing_input(output_dir):
    result = ppt_to_pdf(_request("ppt-to-pdf", inputs=["/no/such/deck.pptx"], output_dir=output_dir))

    assert result.status == "error"
    assert result.error.code == "INPUT_NOT_FOUND"
